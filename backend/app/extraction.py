"""Parse a canonical .xlsx workbook into the same contract as manual JSON."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from .schemas import FinancialInput, LoanTerms

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100


class ExtractionError(ValueError):
    pass


def parse_manual(payload: dict) -> FinancialInput:
    return FinancialInput.model_validate(payload)


def parse_excel(content: bytes, filename: str) -> FinancialInput:
    if Path(filename).suffix.lower() != ".xlsx":
        raise ExtractionError("Only .xlsx files are supported; convert legacy .xls first")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ExtractionError("Workbook must be non-empty and no larger than 5 MiB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ExtractionError("Workbook expanded size exceeds the import limit")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable .xlsx workbook") from exc
    try:
        if "Financials" not in workbook.sheetnames:
            raise ExtractionError("Required sheet: Financials")
        sheet = workbook["Financials"]
        if sheet.max_column is not None and sheet.max_column > 2:
            raise ExtractionError("Only columns A and B are supported")
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=3)
        header = next(rows, ())
        if not header or [cell.value for cell in header[:2]] != ["field", "value"]:
            raise ExtractionError("Row 1 must contain field,value in columns A and B")
        if header[2].value is not None:
            raise ExtractionError("Only columns A and B are supported")
        allowed = set(FinancialInput.model_fields) - {"loan"}
        loan_fields = {f"loan.{field}" for field in LoanTerms.model_fields}
        payload: dict = {}
        seen: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            if all(cell.value is None for cell in row):
                continue
            if row_number > MAX_ROWS + 1:
                raise ExtractionError("Workbook supports at most 100 data rows")
            if any(cell.data_type in {"f", "e"} for cell in row):
                raise ExtractionError(f"Row {row_number}: formulas and Excel errors are not allowed; paste values")
            key, value, extra = (cell.value for cell in row)
            if extra is not None:
                raise ExtractionError(f"Row {row_number}: only columns A and B are supported")
            if not isinstance(key, str) or key.strip() not in allowed | loan_fields:
                raise ExtractionError(f"Row {row_number}: unknown field {key!r}")
            key = key.strip()
            if key in seen:
                raise ExtractionError(f"Row {row_number}: duplicate field {key}")
            seen.add(key)
            if value is None:
                continue
            if key in {"period_start", "period_end"} and hasattr(value, "date"):
                value = value.date()
            if key.startswith("loan."):
                payload.setdefault("loan", {})[key.split(".", 1)[1]] = value
            else:
                payload[key] = value
        return parse_manual(payload)
    except (ExtractionError, ValidationError):
        raise
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable Financials worksheet") from exc
    finally:
        workbook.close()
