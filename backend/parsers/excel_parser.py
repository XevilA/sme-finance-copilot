"""Read English/Thai field-value statements without executing formulas."""

import re
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas import FinancialInput, LoanTerms

DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "Financial_Statement_Translated.xlsx"
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100


class ExtractionError(ValueError):
    """An unsupported or unsafe workbook structure."""


def _normalize(label: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", label).strip().casefold().replace("_", " ").split())


_TRANSLATIONS = {
    "currency": ("สกุลเงิน",),
    "kind": ("ประเภทข้อมูล",),
    "period_start": ("Period Start", "วันเริ่มงวด", "วันที่เริ่มต้น"),
    "period_end": ("Period End", "วันสิ้นงวด", "วันที่สิ้นสุด"),
    "assets": ("Total Assets", "สินทรัพย์", "สินทรัพย์รวม"),
    "liabilities": ("Total Liabilities", "หนี้สิน", "หนี้สินรวม"),
    "equity": ("Total Equity", "ส่วนของเจ้าของ", "ส่วนของผู้ถือหุ้น"),
    "opening_assets": ("Opening Assets", "สินทรัพย์ต้นงวด"),
    "opening_equity": ("Opening Equity", "ส่วนของเจ้าของต้นงวด"),
    "current_assets": ("Current Assets", "สินทรัพย์หมุนเวียน"),
    "current_liabilities": ("Current Liabilities", "หนี้สินหมุนเวียน"),
    "cash_balance": ("Cash Reserve", "Cash Balance", "เงินสดคงเหลือ", "เงินสดสำรอง"),
    "revenue": ("Revenues", "Total Revenue", "รายได้", "รายได้รวม"),
    "total_expenses": ("Expenses", "Total Expenses", "ค่าใช้จ่าย", "ค่าใช้จ่ายรวม"),
    "cost_of_goods_sold": ("COGS", "ต้นทุนขาย"),
    "operating_expenses": ("Operating Expenses", "ค่าใช้จ่ายในการดำเนินงาน"),
    "interest_expense": ("Interest Expense", "ดอกเบี้ยจ่าย"),
    "tax_expense": ("Tax Expense", "ภาษีเงินได้"),
    "net_income": ("Net Profit", "กำไรสุทธิ"),
    "operating_cash_flow": ("Operating Cash Flow", "กระแสเงินสดจากการดำเนินงาน"),
    "investing_cash_flow": ("Investing Cash Flow", "กระแสเงินสดจากการลงทุน"),
    "financing_cash_flow": ("Financing Cash Flow", "กระแสเงินสดจากการจัดหาเงิน"),
    "capital_expenditure": ("CAPEX", "รายจ่ายลงทุน"),
    "cash_available_for_debt_service": ("CFADS", "เงินสดสำหรับชำระหนี้"),
    "debt_principal_due": ("Principal Due", "เงินต้นที่ต้องชำระ"),
    "debt_interest_due": ("Interest Due", "ดอกเบี้ยที่ต้องชำระ"),
    "operating_cash_outflow": ("Operating Cash Outflow", "เงินสดจ่ายจากการดำเนินงาน"),
    "working_capital_needed": ("Working Capital Needed", "เงินทุนหมุนเวียนที่ต้องการ"),
    "minimum_cash_reserve": ("Minimum Cash Reserve", "เงินสดสำรองขั้นต่ำ"),
    "buyer_deposit": ("Buyer Deposit", "เงินมัดจำผู้ซื้อ"),
    "committed_funding": ("Committed Funding", "เงินทุนที่ยืนยันแล้ว"),
}
_FIELDS = (set(FinancialInput.model_fields) - {"loan"}) | {f"loan.{field}" for field in LoanTerms.model_fields}
_ALIASES = {_normalize(field): field for field in _FIELDS}
for _field, _labels in _TRANSLATIONS.items():
    for _label in _labels:
        _ALIASES[_normalize(_label)] = _field

_MONEY_FIELDS = set(FinancialInput.model_fields) - {"currency", "kind", "period_start", "period_end", "loan", "units_sold"}
_GROUPED_AMOUNT = re.compile(r"-?\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?\Z", re.ASCII)
_FIELD_HEADERS = {_normalize(label) for label in ("field", "รายการ")}
_VALUE_HEADERS = {_normalize(label) for label in ("value", "มูลค่า", "จำนวนเงิน")}


def _cell_value(key: str, value: object) -> object:
    if key in {"period_start", "period_end"} and isinstance(value, datetime):
        return value.date()
    if isinstance(value, str) and (key in _MONEY_FIELDS or key == "loan.principal"):
        text = value.strip()
        if text.startswith("(") and text.endswith(")"):
            inner = text[1:-1]
            if not re.fullmatch(r"\d+(?:\.\d{1,2})?", inner, flags=re.ASCII) and not _GROUPED_AMOUNT.fullmatch(inner):
                raise ExtractionError("Invalid accounting amount")
            text = "-" + inner
        if "," in text:
            if not _GROUPED_AMOUNT.fullmatch(text):
                raise ExtractionError("Invalid thousands separators")
            text = text.replace(",", "")
        return text
    return value


def parse_excel(content: bytes, filename: str, *, sheet_name: str = "Financials") -> FinancialInput:
    if Path(filename).suffix.lower() != ".xlsx":
        raise ExtractionError("Only .xlsx files are supported; convert legacy .xls first")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ExtractionError("Workbook must be non-empty and no larger than 5 MiB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ExtractionError("Workbook expanded size exceeds the import limit")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable .xlsx workbook") from exc
    try:
        if sheet_name not in workbook.sheetnames:
            raise ExtractionError(f"Required sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        if sheet.max_column is not None and sheet.max_column > 2:
            raise ExtractionError("Only columns A and B are supported")
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=3)
        header = next(rows, ())
        labels = [_normalize(cell.value) if isinstance(cell.value, str) else None for cell in header[:2]]
        if len(header) != 3 or labels[0] not in _FIELD_HEADERS or labels[1] not in _VALUE_HEADERS:
            raise ExtractionError("Row 1 must contain field,value or their Thai translations in columns A and B")
        if header[2].value is not None:
            raise ExtractionError("Only columns A and B are supported")
        payload: dict = {}
        seen: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            if all(cell.value is None for cell in row):
                continue
            if row_number > MAX_ROWS + 1:
                raise ExtractionError("Workbook supports at most 100 data rows")
            if any(cell.data_type in {"f", "e"} for cell in row):
                raise ExtractionError(f"Row {row_number}: formulas and Excel errors are not allowed; paste values")
            label, value, extra = (cell.value for cell in row)
            if extra is not None:
                raise ExtractionError(f"Row {row_number}: only columns A and B are supported")
            key = _ALIASES.get(_normalize(label)) if isinstance(label, str) else None
            if key is None:
                raise ExtractionError(f"Row {row_number}: unknown field")
            if key in seen:
                raise ExtractionError(f"Row {row_number}: duplicate field {key}")
            seen.add(key)
            if value is None:
                continue
            value = _cell_value(key, value)
            if key.startswith("loan."):
                payload.setdefault("loan", {})[key.split(".", 1)[1]] = value
            else:
                payload[key] = value
        return FinancialInput.model_validate(payload)
    except (ExtractionError, ValidationError):
        raise
    except Exception as exc:
        raise ExtractionError("Invalid or unreadable financial worksheet") from exc
    finally:
        workbook.close()


def parse_excel_file(path: str | Path = DEFAULT_TEMPLATE_PATH, *, sheet_name: str = "Financials") -> FinancialInput:
    file_path = Path(path)
    with file_path.open("rb") as source:
        content = source.read(MAX_UPLOAD_BYTES + 1)
    return parse_excel(content, file_path.name, sheet_name=sheet_name)
